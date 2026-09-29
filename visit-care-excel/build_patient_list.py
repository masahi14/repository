"""訪問診療患者一覧（患者一覧シートのみ・A〜V列のフォーマット）を生成する。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "訪問診療患者一覧.xlsx"
LAST_ROW = 501  # 2〜501行目（500人分）に書式・入力規則・数式を用意

INPUT_FILL = PatternFill("solid", fgColor="FF2F5C8A")    # 手入力列の見出し（青）
FORMULA_FILL = PatternFill("solid", fgColor="FF1F7A6C")  # 自動計算列の見出し（緑）
THIN = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TREATMENT_DH = ["初診", "相談", "SRP", "再評価", "口腔機能", "定期検診",
                "低下症のみ", "無歯顎・低下症のみ", "それ以外"]

# (列, 見出し, 幅, 種類, プルダウン選択肢)
COLUMNS = [
    ("A", "患者ID", 9, "text", None),
    ("B", "名前", 14, "text", None),
    ("C", "介護保険", 10, "list", ["未利用", "利用中"]),
    ("D", "生活保護", 10, "list", ["該当", "非該当"]),
    ("E", "資格確認書", 11, "list", ["あり", "なし"]),
    ("F", "入所施設", 20, "list", ["ハートワン潮来", "養護老人ホーム潮来"]),
    ("G", "訪問頻度", 11, "list", ["月1回", "2か月1回", "不定期"]),
    ("H", "次回目安日", 12, "formula", None),
    ("I", "ドクター最終来院日", 12, "date", None),
    ("J", "ドクター治療内容", 12, "list",
     ["初診", "相談", "EXT", "CR", "義歯印象", "義歯バイト", "義歯試適",
      "義歯セット", "義歯調整", "糸抜き", "その他"]),
    ("K", "ドクター次回予定", 22, "list",
     ["ドクター診察・治療あり", "相談", "ドクター診察・治療なし"]),
    ("L", "EXT", 8, "fraction", None),
    ("M", "CR", 8, "fraction", None),
    ("N", "義歯", 9, "list", ["上のみ", "下のみ", "上下"]),
    ("O", "その他", 12, "text", None),
    ("P", "衛生士最終来院日", 12, "date", None),
    ("Q", "衛生士治療内容", 17, "list", TREATMENT_DH),
    ("R", "衛生士次回予定", 17, "list", TREATMENT_DH),
    ("S", "SRP", 8, "fraction", None),
    ("T", "口腔機能低下症", 12, "date", None),
    ("U", "状態", 8, "list", ["終了"]),
    ("V", "備考", 36, "wrap", None),
]

# 基準日 = I列(ドクター)とP列(衛生士)の新しいほう。月1回→1か月後の同日、2か月1回→2か月後の同日。
H_FORMULA = (
    '=IF(AND($A{r}="",$B{r}=""),"",'
    'IF(MAX($I{r},$P{r})=0,"未定",'
    'IF($G{r}="月1回",EDATE(MAX($I{r},$P{r}),1),'
    'IF($G{r}="2か月1回",EDATE(MAX($I{r},$P{r}),2),"未定"))))'
)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "患者一覧"

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    body_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for col, title, width, kind, choices in COLUMNS:
        ws.column_dimensions[col].width = width
        h = ws[f"{col}1"]
        h.value = title
        h.font = header_font
        h.fill = FORMULA_FILL if kind == "formula" else INPUT_FILL
        h.alignment = center
        h.border = BORDER

        rng = f"{col}2:{col}{LAST_ROW}"
        if kind == "list":
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"',
                                allow_blank=True, showErrorMessage=True,
                                errorTitle="入力エラー",
                                error="プルダウンの選択肢から選んでください。")
            dv.add(rng)
            ws.add_data_validation(dv)
        elif kind == "date":
            dv = DataValidation(type="date", operator="between",
                                formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                                allow_blank=True, showErrorMessage=True,
                                errorTitle="日付を入力してください",
                                error="日付は 2026/8/10 や 8/10 の形で入力してください。"
                                      "（2026.08.10 の形は日付として認識されません）")
            dv.add(rng)
            ws.add_data_validation(dv)

        for r in range(2, LAST_ROW + 1):
            c = ws[f"{col}{r}"]
            c.font = body_font
            c.border = BORDER
            if kind == "wrap":
                c.alignment = Alignment(vertical="center", wrap_text=True)
            elif kind == "text" and col in ("B", "O"):
                c.alignment = Alignment(vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
            if kind in ("date", "formula"):
                c.number_format = "yyyy/mm/dd"
            elif kind == "fraction":
                c.number_format = "@"  # 「1/3」が日付に化けないよう文字列書式
            if kind == "formula":
                c.value = H_FORMULA.format(r=r)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:V{LAST_ROW}"
    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
