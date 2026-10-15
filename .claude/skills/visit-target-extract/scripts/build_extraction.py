#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""訪問患者管理Excelの抽出先シート(歯科衛生士治療 / ドクター治療)を、指定した基準日で更新する。

方針
  * 元ファイルは上書きしない。別ファイル(--out)に書き出す。
  * 触るのは次の部品だけ(他のシート・入力規則・条件付き書式の拡張部分はバイト単位で不変)。
      - 抽出先2シートのXML
      - styles.xml (書式を末尾に追記するだけ)
      - workbook.xml (再計算フラグ・フィルターの名前定義)
      - calcChain.xml (数式を変えるため削除。Excelが自動再作成する)
  * openpyxl で保存し直さない(x14形式のプルダウン/条件付き書式が消えるため)。
    ZIP内のXMLを直接書き換える。
  * 抽出は数式で行う(基準日セルを変えると一覧が再計算される)。
    LibreOfficeで一度再計算し、その結果を「保存済みの値」として埋め込む
    (保護ビュー等、再計算されない環境でも結果が見えるようにするため)。

使い方
  python build_extraction.py inspect --src 元.xlsx
  python build_extraction.py build   --src 元.xlsx --base-date 2026-11-01 --out 出力.xlsx
"""
import argparse
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import warnings
import zipfile
from xml.sax.saxutils import escape

warnings.simplefilter("ignore")
import openpyxl  # noqa: E402  (読み取り専用で使用。保存には使わない)
from lxml import etree  # noqa: E402

# ----------------------------------------------------------------- 設定 ----
SRC_SHEET = "患者一覧"
DR_SHEET = "ドクター治療"
DH_SHEET = "歯科衛生士治療"
MASTER = "マスタ"

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}

# 患者一覧の見出し → 内部名。「同じ意味の見出し」を列挙する(列番号は決め打ちしない)。
HEADING_SYNONYMS = {
    "id": ["患者ID"],
    "name": ["名前", "患者名"],
    "ins": ["介護保険"],
    "fac": ["入所施設"],
    "freq": ["訪問頻度"],
    "next": ["次回目安日"],
    "docLast": ["ドクター最終来院日", "ドクター最終診療日"],
    "docTx": ["ドクター治療内容", "ドクター前回治療内容"],
    "docPlan": ["ドクター次回予定", "ドクター次回予定内容"],
    "hygLast": ["衛生士最終来院日", "衛生士最終診療日"],
    "hygPlan": ["衛生士次回予定", "衛生士次回予定内容"],
    "status": ["状態"],
    "cr": ["CR"],
    "ext": ["抜歯"],
    "den": ["義歯"],
    "srp": ["SRP"],
    "exam": ["低下症最終検査日", "口腔機能低下症の最終検査日", "口腔機能低下症最終検査日"],
    "other": ["その他"],
}
NEEDED = ["id", "name", "ins", "fac", "next", "docLast", "docTx", "docPlan", "hygLast",
          "hygPlan", "status", "cr", "ext", "den", "srp", "exam", "other"]

EXAM_MONTHS = 6        # ドクター側の除外条件②: 口腔機能低下症の最終検査日から何か月経過していないか
ENDED = "終了"          # 状態列の「終了」(既存の抽出条件。マスタの状態リストで確認する)
DATE_MIN = (2000, 1, 1)  # これより前は「無効な日付」(例: 1900/01/02)
DATE_MAX = (2100, 12, 31)

HYG_NOTE = "基準日時点の候補者です。訪問日決定時に、ドクター最終診療日からの期限内か再確認してください。"
DR_NOTE = ("基準日時点のドクター訪問対象患者です。「ドクター診察・治療なし」かつ口腔機能低下症の最終検査日から"
           "6か月が経過していない患者は除外しています(確認が必要な患者は除外せず「要確認」として表示)。")

HELPER_FIRST_COL = 16   # P列。P=抽出キー Q=患者一覧の行番号 R以降=判定用の補助列(すべて非表示)
HELPER_LAST_COL = 30    # AD列


def col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def norm(s):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(s or "")))


class Stop(Exception):
    """質問・確認が必要なため処理を止める(メッセージはそのままユーザーに見せる)。"""


# ------------------------------------------------------- パッケージ操作 ----
def read_zip(path):
    with zipfile.ZipFile(path) as zf:
        return [(i, zf.read(i.filename)) for i in zf.infolist()]


def parts_dict(entries):
    return {i.filename: b for i, b in entries}


def sheet_part_map(parts):
    wb = etree.fromstring(parts["xl/workbook.xml"])
    rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    rid2t = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for idx, s in enumerate(wb.find("m:sheets", NS)):
        t = rid2t[s.get("{%s}id" % NS["r"])]
        out[s.get("name")] = (("xl/" + t) if not t.startswith("/") else t[1:], idx)
    return out


# -------------------------------------------------------- 事前確認(読取) ----
def resolve_headings(ws):
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    heads = {col_letter(i + 1): norm(v) for i, v in enumerate(row) if v not in (None, "")}
    out, problems = {}, []
    for key, syns in HEADING_SYNONYMS.items():
        hits = [c for c, h in heads.items() if h in {norm(s) for s in syns}]
        if len(hits) == 1:
            out[key] = hits[0]
        elif len(hits) > 1:
            problems.append(f"見出し「{syns[0]}」に当たる列が複数あります: {hits}")
        elif key in NEEDED:
            problems.append(f"見出し「{syns[0]}」(同義: {syns[1:] or 'なし'})に当たる列が見つかりません")
    return out, heads, problems


def read_master_list(ws, header):
    """マスタの1行目から見出し(完全一致)を探し、その下の連続した選択肢を返す。"""
    for c in ws[1]:
        if c.value is not None and norm(c.value) == norm(header):
            letter = c.column_letter
            items, r = [], 2
            while ws[f"{letter}{r}"].value not in (None, ""):
                items.append(str(ws[f"{letter}{r}"].value))
                r += 1
            return letter, items
    return None, []


def inspect_source(src):
    """作業前の確認(シート構成・見出し・入力規則・数式・手入力セル)をdictで返す。"""
    entries = read_zip(src)
    parts = parts_dict(entries)
    smap = sheet_part_map(parts)
    info = {"sheets": list(smap), "problems": [], "questions": []}
    for need in (SRC_SHEET, DR_SHEET, DH_SHEET, MASTER):
        if need not in smap:
            info["problems"].append(f"シート「{need}」がありません(シート名が変わった可能性)。")
    if info["problems"]:
        return info
    wb = openpyxl.load_workbook(src, read_only=False, data_only=True)
    cols, heads, probs = resolve_headings(wb[SRC_SHEET])
    info["columns"] = cols
    info["all_headings"] = heads
    info["problems"] += probs
    master = {}
    wsm = wb[MASTER]
    for key, header in (("ins", "介護保険"), ("docTx", "ドクター治療内容"), ("hygTx", "衛生士治療内容"),
                        ("status", "状態"), ("freq", "頻度"), ("freqDays", "次回目安日までの日数")):
        letter, items = read_master_list(wsm, header)
        master[key] = {"col": letter, "items": items}
    info["master"] = master
    # 入力規則(x14拡張を含む): 件数だけ確認する
    dv = {}
    for name, (part, _) in smap.items():
        x = parts[part].decode("utf8")
        dv[name] = {"standard": len(re.findall(r"<dataValidation[ >]", x)),
                    "x14": len(re.findall(r"<x14:dataValidation[ >]", x)),
                    "x14_cf": len(re.findall(r"<x14:cfRule[ >]", x))}
    info["validations"] = dv
    info["protected"] = {n: bool(re.search(r"<sheetProtection[^>]*sheet=\"1\"", parts[p[0]].decode("utf8")))
                         for n, p in smap.items()}
    info["macro"] = any(i.filename.endswith("vbaProject.bin") for i, _ in entries)
    # シート間の連動: 抽出先2シートを参照している他の部品(他シートの数式・名前定義)と、抽出先が参照しているシート
    refs_to, refs_from = [], {}
    for name, (part, _) in smap.items():
        x = parts[part].decode("utf8")
        if name not in (DR_SHEET, DH_SHEET):
            for t in (DR_SHEET, DH_SHEET):
                if re.search(rf"<f[^>]*>[^<]*{t}!", x):
                    refs_to.append(f"{name} → {t}")
        else:
            refs_from[name] = sorted(set(re.findall(r"([^\s!(,=&<>\"+*/-]+)!\$", x)) - {name})
    wbx = parts["xl/workbook.xml"].decode("utf8")
    for t in (DR_SHEET, DH_SHEET):
        for m in re.finditer(rf'<definedName name="([^"]+)"[^>]*>[^<]*{t}!', wbx):
            refs_to.append(f"名前定義 {m.group(1)} → {t}")
    info["links"] = {"others_refer_to_extract_sheets": refs_to, "extract_sheets_refer_to": refs_from}
    info["manual_cells"] = find_manual_cells(parts, smap)
    # 抽出先の既存の抽出条件(P列の数式)
    for sheet in (DR_SHEET, DH_SHEET):
        m = re.search(r'<c r="P4"[^>]*><f>(.*?)</f>', parts[smap[sheet][0]].decode("utf8"))
        info.setdefault("existing_filter", {})[sheet] = (
            re.sub(r"&quot;", '"', re.sub(r"&amp;", "&", m.group(1))) if m else None)
    return info


def find_manual_cells(parts, smap):
    """抽出先2シートにある「数式でも見出しでもない」セル=手入力の予定・メモ候補を返す。"""
    found = {}
    for sheet in (DR_SHEET, DH_SHEET):
        root = etree.fromstring(parts[smap[sheet][0]])
        styles = parts["xl/styles.xml"].decode("utf8")
        unlocked = unlocked_style_ids(styles)
        cells = []
        for c in root.iter("{%s}c" % NS["m"]):
            ref = c.get("r")
            col = col_index(re.match(r"[A-Z]+", ref).group(0))
            row = int(re.search(r"\d+", ref).group(0))
            has_f = c.find("m:f", NS) is not None
            v = c.find("m:v", NS)
            isel = c.find("m:is", NS)
            has_val = (v is not None and (v.text or "") != "") or isel is not None
            if has_f:
                continue
            # 既知の固定セル: 1行目のタイトル/補助見出し、2行目の基準日、3行目の見出し
            if row == 1 and (col == 1 or col >= HELPER_FIRST_COL):
                continue
            if row == 2 and col in (1, 2):
                continue
            if row == 3 and col <= 13:
                continue
            if has_val or (c.get("s") and int(c.get("s")) in unlocked):
                cells.append(ref)
        if cells:
            found[sheet] = cells
    return found


def unlocked_style_ids(styles_xml):
    m = re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles_xml, re.S)
    xfs = re.findall(r"<xf\b.*?(?:/>|</xf>)", m.group(1), re.S) if m else []
    return {i for i, x in enumerate(xfs) if 'locked="0"' in x}


# ------------------------------------------------------------ 書式の追記 ----
class Styles:
    """styles.xml に、同一定義があれば再利用し、なければ末尾へ追記する。"""

    def __init__(self, xml):
        self.xml = xml

    def _block(self, tag):
        m = re.search(rf"<{tag}\b[^>]*>(.*?)</{tag}>", self.xml, re.S)
        return m

    def _add(self, tag, item_re, new_item):
        m = self._block(tag)
        items = re.findall(item_re, m.group(1), re.S)
        if new_item in items:
            return items.index(new_item)
        count = len(items)
        head = re.search(rf"<{tag}\b[^>]*>", self.xml).group(0)
        new_head = re.sub(r'count="\d+"', f'count="{count + 1}"', head)
        self.xml = self.xml.replace(head, new_head, 1)
        self.xml = self.xml.replace(f"</{tag}>", new_item + f"</{tag}>", 1)
        return count

    def font(self, x):
        return self._add("fonts", r"<font>.*?</font>", x)

    def fill(self, x):
        return self._add("fills", r"<fill>.*?</fill>", x)

    def xf(self, x):
        return self._add("cellXfs", r"<xf\b[^>]*?(?:/>|>.*?</xf>)", x)


def prepare_styles(styles_xml):
    st = Styles(styles_xml)
    f_bold = st.font('<font><b/><sz val="10"/><name val="Arial"/><family val="2"/></font>')
    f_note = st.font('<font><b/><sz val="10"/><color rgb="FF9C0006"/><name val="Arial"/><family val="2"/></font>')
    fill_in = st.fill('<fill><patternFill patternType="solid"><fgColor rgb="FFFFF2CC"/></patternFill></fill>')
    # 既存のセルxf(26=本文, 49=日付)を踏襲し、同じ罫線(borderId=3)を使う
    ids = {}
    ids["label"] = st.xf(f'<xf numFmtId="0" fontId="{f_bold}" fillId="0" borderId="0" xfId="0" applyFont="1" '
                         f'applyAlignment="1"><alignment horizontal="left" vertical="center"/></xf>')
    ids["date_in"] = st.xf(f'<xf numFmtId="176" fontId="{f_bold}" fillId="{fill_in}" borderId="3" xfId="0" '
                           f'applyNumberFormat="1" applyFont="1" applyFill="1" applyBorder="1" applyAlignment="1" '
                           f'applyProtection="1"><alignment horizontal="center" vertical="center"/>'
                           f'<protection locked="0"/></xf>')
    ids["date_ref"] = st.xf(f'<xf numFmtId="176" fontId="{f_bold}" fillId="0" borderId="3" xfId="0" '
                            f'applyNumberFormat="1" applyFont="1" applyBorder="1" applyAlignment="1">'
                            f'<alignment horizontal="center" vertical="center"/></xf>')
    ids["note"] = st.xf(f'<xf numFmtId="0" fontId="{f_note}" fillId="0" borderId="0" xfId="0" applyFont="1" '
                        f'applyAlignment="1"><alignment horizontal="left" vertical="center"/></xf>')
    ids["reason"] = st.xf('<xf numFmtId="0" fontId="15" fillId="0" borderId="3" xfId="0" applyFont="1" '
                          'applyBorder="1" applyAlignment="1"><alignment horizontal="left" vertical="center" wrapText="1"/></xf>')
    ids["summary"] = st.xf(f'<xf numFmtId="0" fontId="{f_bold}" fillId="0" borderId="0" xfId="0" applyFont="1" '
                           f'applyAlignment="1"><alignment horizontal="left" vertical="center"/></xf>')
    return st.xml, ids


# ----------------------------------------------------------- 数式の生成 ----
class Ctx:
    """抽出先シート1枚分の数式を作る共通部品。"""

    def __init__(self, cols, master_ref, last_row):
        self.c = cols                  # 内部名 → 患者一覧の列文字
        self.m = master_ref            # マスタのセル参照
        self.last = last_row           # 患者一覧の最終行(補助列の最終行)
        self.slots = last_row - 1      # 表示枠の行数
        self.first_slot_row = 4
        self.last_slot_row = 3 + self.slots

    def pl(self, key, n):              # 患者一覧!$X{n}
        return f"{SRC_SHEET}!${self.c[key]}{n}"

    def plcol(self, key):              # 患者一覧!$X:$X
        return f"{SRC_SHEET}!${self.c[key]}:${self.c[key]}"

    @staticmethod
    def txt(x):
        return f'IF(ISNUMBER({x}),TEXT({x},"yyyy/mm/dd"),IF(TRIM({x}&"")="","空欄",{x}&""))'


def date_state(ref, base="$B$2"):
    """空欄/無効/基準日より後/正常 の判定式(文字列を返す)。"""
    y0, m0, d0 = DATE_MIN
    y1, m1, d1 = DATE_MAX
    return (f'IF(TRIM({ref}&"")="","blank",IF(NOT(ISNUMBER({ref})),"invalid",'
            f'IF(OR({ref}<DATE({y0},{m0},{d0}),{ref}>DATE({y1},{m1},{d1})),"invalid",'
            f'IF({ref}>{base},"future","ok"))))')


def hyg_helpers(cx, n):
    """歯科衛生士治療シートの補助列(患者一覧のn行目に対応)。"""
    M = cx.m
    I, L, Uc = cx.pl("docLast", n), cx.pl("hygLast", n), cx.pl("exam", n)
    C, N = cx.pl("ins", n), cx.pl("hygPlan", n)
    f = {}
    f["R"] = f'IF(OR({cx.pl("id", n)}="",{cx.pl("status", n)}="{ENDED}"),0,1)'
    f["S"] = date_state(I)
    f["T"] = f'IF(AND(ISNUMBER({L}),{L}>$B$2),1,0)'
    f["U"] = f'IF(AND(ISNUMBER({Uc}),{Uc}>$B$2),1,0)'
    f["V"] = f'IF(TRIM({C}&"")={M["ins_use"]},"U",IF(TRIM({C}&"")={M["ins_not"]},"N","?"))'
    f["W"] = (f'IF(TRIM({N}&"")="","?",IF(TRIM({N}&"")={M["hyg_regular"]},"T",'
              f'IF(ISNUMBER(MATCH(TRIM({N}&""),{M["hyg_list"]},0)),"O","?")))')
    f["X"] = f'IF($V{n}="U",3,IF($W{n}="T",2,1))'
    f["Y"] = f'IF($V{n}="U",3,IF($V{n}="?",3,IF($W{n}="T",2,IF($W{n}="O",1,2))))'
    f["Z"] = f'IF(COUNTIF({SRC_SHEET}!${cx.c["id"]}$2:${cx.c["id"]}${cx.last},{cx.pl("id", n)})>1,1,0)'
    f["AA"] = f'IF($S{n}="ok",IF($B$2<=EDATE({I},$X{n}),1,0),0)'
    f["AB"] = f'IF($S{n}="ok",IF($B$2<=EDATE({I},$Y{n}),1,0),0)'
    f["AC"] = (f'IF(OR($R{n}=0,NOT(ISNUMBER($B$2))),"",IF(OR($S{n}<>"ok",$AB{n}<>$AA{n},$T{n}=1,$U{n}=1,$Z{n}=1),'
               f'"C",IF($AA{n}=1,"T","")))')
    t = cx.txt
    seg_doc = (f'IF($S{n}="blank","ドクター最終診療日が空欄のため、期限内か判定できません。",'
               f'IF($S{n}="invalid","ドクター最終診療日（"&{t(I)}&"）が有効な日付ではないため、判定できません。",'
               f'IF($S{n}="future","ドクター最終診療日（"&{t(I)}&"）が基準日より後のため、判定に使用していません。","")))')
    seg_ins = (f'IF(AND($V{n}="?",OR($S{n}<>"ok",$AB{n}<>$AA{n})),'
               f'"介護保険の利用状況が空欄または不明（"&{t(C)}&"）のため、3か月・2か月・1か月のどの期限か決められません。","")')
    seg_plan = (f'IF(AND($V{n}<>"U",$W{n}="?",OR($S{n}<>"ok",$AB{n}<>$AA{n})),'
                f'"衛生士次回予定が空欄、または「定期検診」と同じ意味か判断できない表記（"&{t(N)}&"）のため、2か月か1か月か決められません。","")')
    seg_hyg = f'IF($T{n}=1,"衛生士最終診療日（"&{t(L)}&"）が基準日より後です。","")'
    seg_exam = f'IF($U{n}=1,"口腔機能低下症の最終検査日（"&{t(Uc)}&"）が基準日より後です。","")'
    seg_dup = f'IF($Z{n}=1,"患者IDが重複しており、患者情報の対応を特定できません。","")'
    f["AD"] = (f'IF($AC{n}<>"C","","要確認："&{seg_doc}&{seg_ins}&{seg_plan}&{seg_hyg}&{seg_exam}&{seg_dup})')
    f["P"] = (f'IF(OR({cx.pl("id", n)}="",{cx.pl("status", n)}="{ENDED}"),"",'
              f'IF($AC{n}="T",{n},IF($AC{n}="C",{n}+1000,"")))')
    return f


def dr_helpers(cx, n):
    """ドクター治療シートの補助列。"""
    M = cx.m
    Uc, K = cx.pl("exam", n), cx.pl("docPlan", n)
    f = {}
    f["R"] = f'IF(OR({cx.pl("id", n)}="",{cx.pl("status", n)}="{ENDED}"),0,1)'
    f["S"] = date_state(Uc)
    f["T"] = (f'IF(TRIM({K}&"")="","?",IF(TRIM({K}&"")={M["doc_none"]},"Y",'
              f'IF(ISNUMBER(MATCH(TRIM({K}&""),{M["doc_list"]},0)),"N","?")))')
    f["U"] = f'IF($S{n}="ok",IF($B$2<EDATE({Uc},{EXAM_MONTHS}),"T","F"),"?")'
    f["V"] = (f'IF(OR({cx.pl("cr", n)}&""<>"",{cx.pl("ext", n)}&""<>"",{cx.pl("den", n)}&""<>""),1,0)')
    f["W"] = f'IF(COUNTIF({SRC_SHEET}!${cx.c["id"]}$2:${cx.c["id"]}${cx.last},{cx.pl("id", n)})>1,1,0)'
    f["X"] = f'IF(AND($T{n}="Y",$U{n}="T"),1,0)'
    f["Y"] = f'IF(OR($T{n}="N",$U{n}="F"),1,0)'
    f["Z"] = f'IF(AND($X{n}=0,$Y{n}=0),1,0)'
    f["AA"] = f'IF(AND($T{n}="Y",$V{n}=1),1,0)'
    f["AB"] = f'IF($S{n}="future",1,0)'
    f["AC"] = (f'IF(OR($R{n}=0,NOT(ISNUMBER($B$2))),"",IF(OR($Z{n}=1,$AA{n}=1,$AB{n}=1,$W{n}=1),"C",'
               f'IF($X{n}=1,"",IF($Y{n}=1,"T",""))))')
    t = cx.txt
    seg_k = (f'IF(AND($Z{n}=1,$T{n}="?"),IF(TRIM({K}&"")="","ドクター次回予定が空欄のため、「ドクター診察・治療なし」か判断できません。",'
             f'"ドクター次回予定の表記（"&{K}&"）が「ドクター診察・治療なし」と同じ意味か判断できません。"),"")')
    seg_e = (f'IF(AND($Z{n}=1,$U{n}="?",$S{n}<>"future"),IF($S{n}="blank",'
             f'"口腔機能低下症の最終検査日が空欄のため、検査から{EXAM_MONTHS}か月が経過していないか判断できません。",'
             f'"口腔機能低下症の最終検査日（"&{t(Uc)}&"）が有効な日付ではないため、検査から{EXAM_MONTHS}か月が経過していないか判断できません。"),"")')
    seg_f = f'IF($S{n}="future","口腔機能低下症の最終検査日（"&{t(Uc)}&"）が基準日より後のため、判定に使用していません。","")'
    seg_c = (f'IF($AA{n}=1,"ドクター次回予定が「ドクター診察・治療なし」ですが、CR・抜歯・義歯に進捗の記録があります（"'
             f'&TRIM(IF({cx.pl("cr", n)}&""<>"","CR:"&{cx.pl("cr", n)}&" ","")'
             f'&IF({cx.pl("ext", n)}&""<>"","抜歯:"&{cx.pl("ext", n)}&" ","")'
             f'&IF({cx.pl("den", n)}&""<>"","義歯:"&{cx.pl("den", n)}&" ",""))'
             f'&"）。進捗表記の意味が未確認のため、未完了か判断できません。","")')
    seg_dup = f'IF($W{n}=1,"患者IDが重複しており、患者情報の対応を特定できません。","")'
    f["AD"] = f'IF($AC{n}<>"C","","要確認："&{seg_k}&{seg_e}&{seg_f}&{seg_c}&{seg_dup})'
    f["P"] = (f'IF(OR({cx.pl("id", n)}="",{cx.pl("status", n)}="{ENDED}"),"",'
              f'IF($AC{n}="T",{n},IF($AC{n}="C",{n}+1000,"")))')
    return f


HYG_HELPER_HEADS = {"P": "抽出キー(確定=行番号/要確認=行番号+1000)", "Q": "患者一覧の行番号", "R": "対象患者(状態≠終了)",
                    "S": "ドクター最終診療日の状態", "T": "衛生士最終診療日>基準日", "U": "検査日>基準日",
                    "V": "介護保険(U利用中/N未利用/?不明)", "W": "衛生士次回予定(T定期検診/O以外/?不明)",
                    "X": "適用月数(最短)", "Y": "適用月数(最長)", "Z": "患者ID重複", "AA": "期限内(最短の月数)",
                    "AB": "期限内(最長の月数)", "AC": "区分(T確定/C要確認)", "AD": "要確認の理由"}
DR_HELPER_HEADS = {"P": "抽出キー(確定=行番号/要確認=行番号+1000)", "Q": "患者一覧の行番号", "R": "対象患者(状態≠終了)",
                   "S": "検査日の状態", "T": "ドクター次回予定(Y治療なし/N他/?不明)", "U": f"検査から{EXAM_MONTHS}か月未満(T/F/?)",
                   "V": "CR・抜歯・義歯に記録あり", "W": "患者ID重複", "X": "除外条件に該当", "Y": "除外されないことが確定",
                   "Z": "判断できない", "AA": "治療なしと進捗記録が矛盾", "AB": "検査日>基準日",
                   "AC": "区分(T確定/C要確認)", "AD": "要確認の理由"}


# ------------------------------------------------------------ セルXML ----
def x_text(ref, s, text):
    return f'<c r="{ref}" s="{s}" t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>'


def x_num(ref, s, v):
    return f'<c r="{ref}" s="{s}"><v>{v}</v></c>'


def x_formula(ref, s, formula, cache):
    s_attr = f' s="{s}"' if s is not None else ""
    f = f"<f>{escape(formula)}</f>"
    if cache is None:
        return f'<c r="{ref}"{s_attr}>{f}</c>'
    t, v = cache.get(ref, ("str", ""))
    if t == "n":
        return f'<c r="{ref}"{s_attr}>{f}<v>{v}</v></c>'
    if t == "b":
        return f'<c r="{ref}"{s_attr} t="b">{f}<v>{v}</v></c>'
    return f'<c r="{ref}"{s_attr} t="str">{f}<v>{escape(v)}</v></c>'


def assemble_rows(cells, row_attrs):
    """cells: {(row,col): xml}。行・列の昇順で並べてsheetData本体を作る。"""
    rows = {}
    for (r, c), x in cells.items():
        rows.setdefault(r, []).append((c, x))
    out = []
    for r in sorted(rows):
        attrs = row_attrs.get(r, "")
        out.append(f'<row r="{r}"{attrs}>' + "".join(x for _, x in sorted(rows[r])) + "</row>")
    return "".join(out)


# ------------------------------------------------ 条件付き書式・拡張部分 ----
def parse_cf_blocks(xml):
    out = []
    for m in re.finditer(r'<conditionalFormatting sqref="([^"]+)">(.*?)</conditionalFormatting>', xml, re.S):
        rules = re.findall(r'<cfRule type="expression" dxfId="(\d+)" priority="(\d+)"><formula>(.*?)</formula></cfRule>',
                           m.group(2), re.S)
        out.append({"sqref": m.group(1), "rules": [(int(a), int(b), c) for a, b, c in rules]})
    return out


def cf_xml(sqref, rules, col):
    body = "".join(
        f'<cfRule type="expression" dxfId="{d}" priority="{p}"><formula>{re.sub(r"[$][A-Z]+4", f"${col}4", f)}</formula></cfRule>'
        for d, p, f in rules)
    return f'<conditionalFormatting sqref="{sqref}">{body}</conditionalFormatting>'


def retarget_x14(ext_xml, col, last_row):
    """施設名の色分け(x14形式)を、施設列が動いた先へ付け替える。"""
    ext_xml = re.sub(r"MATCH\(\$[A-Z]+4,", f"MATCH(${col}4,", ext_xml)
    ext_xml = re.sub(r"<xm:sqref>[A-Z]+4:[A-Z]+\d+</xm:sqref>", f"<xm:sqref>{col}4:{col}{last_row}</xm:sqref>", ext_xml)
    return ext_xml


# --------------------------------------------------- シートXMLの組み立て ----
def split_sheet(xml):
    """元のシートXMLから、そのまま引き継ぐ部品を取り出す。"""
    head = xml[:xml.index("<dimension")]
    keep = {}
    for tag in ("sheetProtection", "phoneticPr", "pageMargins", "pageSetup"):
        m = re.search(rf"<{tag}\b[^>]*/>", xml)
        keep[tag] = m.group(0) if m else ""
    m = re.search(r"<extLst>.*</extLst>", xml, re.S)
    keep["extLst"] = m.group(0) if m else ""
    keep["cf"] = parse_cf_blocks(xml)
    keep["E4"] = None
    m = re.search(r'<c r="E4"[^>]*><f>(.*?)</f>', xml, re.S)
    if m:
        keep["E4"] = m.group(1)
    return head, keep


def cols_xml(widths, hidden_cols):
    out = []
    for i in range(1, HELPER_FIRST_COL):
        w = widths.get(i, 10)
        h = ' hidden="1"' if i in hidden_cols else ""
        out.append(f'<col min="{i}" max="{i}" width="{w}"{h} customWidth="1"/>')
    out.append(f'<col min="{HELPER_FIRST_COL}" max="{HELPER_LAST_COL}" width="13" hidden="1" customWidth="1"/>')
    return "<cols>" + "".join(out) + "</cols>"


def build_sheet(kind, orig_xml, cx, st, base_serial, cache):
    head, keep = split_sheet(orig_xml)
    cells, row_attrs = {}, {1: ' ht="18" customHeight="1"', 2: ' ht="18" customHeight="1"',
                            3: ' ht="25.5" customHeight="1"'}
    is_h = kind == "hyg"
    helpers_fn = hyg_helpers if is_h else dr_helpers
    heads_map = HYG_HELPER_HEADS if is_h else DR_HELPER_HEADS
    hdr_style = 29 if is_h else 27
    TEXT, DATE, DATE_IN = 26, 49, st["date_in"]
    ls, le = cx.first_slot_row, cx.last_slot_row

    # --- 1〜2行目: 注記・基準日・件数 ---
    cells[(1, 1)] = x_text("A1", st["note"], HYG_NOTE if is_h else DR_NOTE)
    cells[(2, 1)] = x_text("A2", st["label"], "抽出基準日")
    if is_h:
        cells[(2, 2)] = x_num("B2", DATE_IN, base_serial)
    else:
        cells[(2, 2)] = x_formula("B2", st["date_ref"], f"{DH_SHEET}!$B$2", cache.get(DR_SHEET) if cache else None)
    sheet_name = DH_SHEET if is_h else DR_SHEET
    sc = cache.get(sheet_name) if cache else None
    label = "衛生士のみで訪問できる候補者" if is_h else "ドクター訪問対象"
    cells[(2, 4)] = x_formula(
        "D2", st["summary"],
        f'IF(NOT(ISNUMBER($B$2)),"基準日を入力してください",'
        f'"{label}：確定 "&COUNTIF($AC$2:$AC${cx.last},"T")&" 名 ／ 要確認 "&COUNTIF($AC$2:$AC${cx.last},"C")&" 名")',
        sc)
    for col, h in heads_map.items():
        cells[(1, col_index(col))] = x_text(f"{col}1", TEXT, h)

    # --- 3行目: 見出し ---
    if is_h:
        layout = [("患者ID", "id", "plain"), ("患者名", "name", "plain"), ("介護保険", "ins", "guard"),
                  ("入所施設", "fac", "guard"), ("次回目安日", "next", "date"),
                  ("ドクター最終診療日", "docLast", "date"), ("衛生士最終診療日", "hygLast", "date"),
                  ("衛生士次回予定", "hygPlan", "guard"), ("SRP", "srp", "guard"),
                  ("口腔機能低下症の最終検査日", "exam", "date"), ("その他", "other", "guard")]
        reason_col = 12
    else:
        layout = [("患者ID", "id", "plain"), ("名前", "name", "plain"), ("入所施設", "fac", "guard"),
                  ("ドクター最終来院日", "docLast", "date"), ("ドクター次回目安日", "E", "special"),
                  ("ドクター治療内容", "docTx", "guard"), ("ドクター次回予定内容", "docPlan", "guard"),
                  ("CR", "cr", "guard"), ("抜歯", "ext", "guard"), ("義歯", "den", "guard"),
                  ("低下症最終検査日", "exam", "date"), ("その他", "other", "guard")]
        reason_col = 13
    for i, (h, _, _) in enumerate(layout, start=1):
        cells[(3, i)] = x_text(f"{col_letter(i)}3", hdr_style, h)
    cells[(3, reason_col)] = x_text(f"{col_letter(reason_col)}3", hdr_style, "要確認（確認が必要な理由）")

    # --- 補助列(患者一覧の2〜最終行と同じ行番号) ---
    for n in range(2, cx.last + 1):
        hf = helpers_fn(cx, n)
        for col, fm in hf.items():
            ref = f"{col}{n}"
            cells[(n, col_index(col))] = x_formula(ref, None, fm, sc)
    # --- 抽出表本体 ---
    for k, r in enumerate(range(ls, le + 1), start=1):
        cells[(r, 17)] = x_formula(f"Q{r}", None, f'IFERROR(MOD(SMALL($P$2:$P${cx.last},{k}),1000),"")', sc)
        for i, (_, key, mode) in enumerate(layout, start=1):
            ref = f"{col_letter(i)}{r}"
            q = f"$Q{r}"
            if mode == "special":
                f0 = keep["E4"]
                if not f0:
                    raise Stop("ドクター治療シートの「次回目安日」列の数式を取得できませんでした。")
                f = re.sub(r"\$Q4\b", q, unescape_f(f0))
                cells[(r, i)] = x_formula(ref, DATE, f, sc)
                continue
            plc = cx.plcol(key)
            if mode == "plain":
                f = f'IF({q}="","",INDEX({plc},{q}))'
                cells[(r, i)] = x_formula(ref, TEXT, f, sc)
            else:
                f = f'IF({q}="","",IF(INDEX({plc},{q})="","",INDEX({plc},{q})))'
                cells[(r, i)] = x_formula(ref, DATE if mode == "date" else TEXT, f, sc)
        rc = col_letter(reason_col)
        cells[(r, reason_col)] = x_formula(f"{rc}{r}", st["reason"],
                                           f'IF({"$Q%d" % r}="","",INDEX($AD:$AD,$Q{r}))', sc)

    sheet_data = "<sheetData>" + assemble_rows(cells, row_attrs) + "</sheetData>"
    widths = ({1: 11, 2: 20, 3: 10, 4: 20, 5: 13, 6: 14, 7: 14, 8: 14, 9: 9, 10: 15, 11: 30, 12: 70}
              if is_h else
              {1: 11, 2: 20, 3: 20, 4: 14, 5: 16, 6: 12, 7: 18, 8: 10, 9: 10, 10: 10, 11: 16, 12: 30, 13: 70})
    hidden = set() if is_h else {4, 6}   # ドクター最終来院日・ドクター治療内容は表示しない(列は削除せず非表示)
    nlast = reason_col
    sv = ('<sheetViews><sheetView showGridLines="0" workbookViewId="0"><pane ySplit="3" topLeftCell="A4" '
          'activePane="bottomLeft" state="frozen"/><selection pane="bottomLeft" activeCell="A4" sqref="A4"/>'
          '</sheetView></sheetViews>')

    # --- 条件付き書式 ---
    blocks = keep["cf"]
    cfx = ""
    if is_h:
        aging = next(b for b in blocks if any("EDATE(TODAY()" in r[2] for r in b["rules"]))
        plans = [b for b in blocks if b is not aging and b["rules"] and re.search(r'="[^"]+"$', b["rules"][0][2])]
        if not plans:
            raise Stop("歯科衛生士治療シートの次回予定の色分けルールが見つかりません。")
        plan = max(plans, key=lambda b: min(p for _, p, _ in b["rules"]))
        cfx += cf_xml(f"E{ls}:E{le} G{ls}:G{le}", aging["rules"], "G")
        cfx += cf_xml(f"H{ls}:H{le}", plan["rules"], "H")
    else:
        for b in blocks:
            sq = re.sub(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", lambda m: f"{m.group(1)}{m.group(2)}:{m.group(3)}{le}", b["sqref"])
            cfx += (f'<conditionalFormatting sqref="{sq}">' + "".join(
                f'<cfRule type="expression" dxfId="{d}" priority="{p}"><formula>{f}</formula></cfRule>'
                for d, p, f in b["rules"]) + "</conditionalFormatting>")
    ext = keep["extLst"]
    if ext:
        ext = retarget_x14(ext, "D" if is_h else "C", le)
    af = f'<autoFilter ref="A3:{col_letter(nlast)}{le}"/>'
    dim = f'<dimension ref="A1:{col_letter(HELPER_LAST_COL)}{le}"/>'
    xml = (head + dim + sv + '<sheetFormatPr defaultRowHeight="13.5"/>' + cols_xml(widths, hidden) + sheet_data
           + keep["sheetProtection"] + af + keep["phoneticPr"] + cfx + keep["pageMargins"] + keep["pageSetup"] + ext
           + "</worksheet>")
    return xml, nlast


def unescape_f(s):
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'")
            .replace("&amp;", "&"))


# ---------------------------------------------------- workbook.xml 等 ----
def patch_workbook(wbxml, dr_idx, dh_idx, dr_last_col, dh_last_col, le):
    if "fullCalcOnLoad" not in wbxml:
        wbxml = re.sub(r"<calcPr\b([^>]*?)/>", lambda m: f'<calcPr{m.group(1)} fullCalcOnLoad="1"/>', wbxml, count=1)
    # 既存の抽出先フィルター名を除いて作り直す(再実行しても重複しない)
    def drop(m):
        return "" if m.group(1) in (str(dr_idx), str(dh_idx)) else m.group(0)
    wbxml = re.sub(r'<definedName name="_xlnm\._FilterDatabase" localSheetId="(\d+)"[^>]*>[^<]*</definedName>', drop, wbxml)
    add = (f'<definedName name="_xlnm._FilterDatabase" localSheetId="{dr_idx}" hidden="1">'
           f'{DR_SHEET}!$A$3:${col_letter(dr_last_col)}${le}</definedName>'
           f'<definedName name="_xlnm._FilterDatabase" localSheetId="{dh_idx}" hidden="1">'
           f'{DH_SHEET}!$A$3:${col_letter(dh_last_col)}${le}</definedName>')
    if "<definedNames>" in wbxml:
        wbxml = wbxml.replace("</definedNames>", add + "</definedNames>", 1)
    else:
        wbxml = wbxml.replace("<calcPr", f"<definedNames>{add}</definedNames><calcPr", 1)
    return wbxml


def drop_calc_chain(parts, order):
    if "xl/calcChain.xml" in parts:
        del parts["xl/calcChain.xml"]
        order[:] = [n for n in order if n != "xl/calcChain.xml"]
        parts["[Content_Types].xml"] = re.sub(
            rb'<Override PartName="/xl/calcChain.xml"[^>]*/>', b"", parts["[Content_Types].xml"])
        parts["xl/_rels/workbook.xml.rels"] = re.sub(
            rb'<Relationship [^>]*Target="calcChain.xml"[^>]*/>', b"", parts["xl/_rels/workbook.xml.rels"])


# ------------------------------------------------------- LibreOffice再計算 ----
def lo_recalc(path, outdir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    prof = tempfile.mkdtemp(prefix="lo_profile_")
    try:
        subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", outdir, path],
                       check=True, capture_output=True, timeout=300)
    except Exception as e:  # noqa: BLE001
        print(f"[warn] LibreOfficeでの再計算に失敗: {e}", file=sys.stderr)
        return None
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    out = os.path.join(outdir, os.path.basename(path))
    return out if os.path.exists(out) else None


def read_cached_values(xlsx, sheet_names):
    """LibreOffice出力から、数式セルの計算値(型, 値)を取り出す。"""
    parts = parts_dict(read_zip(xlsx))
    smap = sheet_part_map(parts)
    sst = []
    if "xl/sharedStrings.xml" in parts:
        root = etree.fromstring(parts["xl/sharedStrings.xml"])
        for si in root.findall("m:si", NS):
            sst.append("".join(t.text or "" for t in si.iter("{%s}t" % NS["m"])))
    res = {}
    for name in sheet_names:
        root = etree.fromstring(parts[smap[name][0]])
        d = {}
        for c in root.iter("{%s}c" % NS["m"]):
            if c.find("m:f", NS) is None:
                continue
            t, v = c.get("t"), c.find("m:v", NS)
            val = "" if v is None or v.text is None else v.text
            if t == "e":
                raise Stop(f"{name}!{c.get('r')} が数式エラー({val})になりました。")
            if t == "s":
                d[c.get("r")] = ("str", sst[int(val)])
            elif t in ("str", "inlineStr"):
                d[c.get("r")] = ("str", val)
            elif t == "b":
                d[c.get("r")] = ("b", val)
            else:
                d[c.get("r")] = ("n", val) if val != "" else ("str", "")
        res[name] = d
    return res


# -------------------------------------------------------------- 本体 ----
def master_refs(info):
    m = info["master"]
    def need(key, label):
        if not m[key]["col"] or not m[key]["items"]:
            raise Stop(f"マスタに「{label}」の選択肢が見つかりません。")
    for k, lab in (("ins", "介護保険"), ("docTx", "ドクター治療内容"), ("hygTx", "衛生士治療内容"), ("status", "状態")):
        need(k, lab)
    def addr(key, value):
        col, items = m[key]["col"], m[key]["items"]
        idx = [i for i, v in enumerate(items) if norm(v) == norm(value)]
        if len(idx) != 1:
            raise Stop(f"マスタ「{key}」の選択肢に「{value}」が1つだけ存在することを確認できません(選択肢: {items})。"
                       f"表記が変わっている場合は、どの選択肢が同じ意味か確認してください。")
        return f"{MASTER}!${col}${idx[0] + 2}"
    def rng(key):
        col, items = m[key]["col"], m[key]["items"]
        return f"{MASTER}!${col}$2:${col}${len(items) + 1}"
    if norm(ENDED) not in [norm(x) for x in m["status"]["items"]]:
        raise Stop(f"マスタの状態の選択肢に「{ENDED}」がありません: {m['status']['items']}")
    return {"ins_use": addr("ins", "利用中"), "ins_not": addr("ins", "未利用"),
            "hyg_regular": addr("hygTx", "定期検診"), "hyg_list": rng("hygTx"),
            "doc_none": addr("docTx", "ドクター診察・治療なし"), "doc_list": rng("docTx")}


def run_build(src, base_date, out, allow_manual=False, no_lo=False):
    info = inspect_source(src)
    if info["problems"]:
        raise Stop("元ファイルの確認で問題が見つかりました:\n- " + "\n- ".join(info["problems"]))
    if info["macro"]:
        print("[info] マクロ付きファイルです。vbaProjectは変更せずに保持します。", file=sys.stderr)
    if info["manual_cells"] and not allow_manual:
        raise Stop("抽出先シートに手入力と思われるセルがあります。患者単位で保持できるか確認してください:\n"
                   + json.dumps(info["manual_cells"], ensure_ascii=False))
    refs = master_refs(info)
    wb = openpyxl.load_workbook(src, read_only=True)
    max_row = wb[SRC_SHEET].max_row or 301
    last = max(301, max_row)
    cx = Ctx(info["columns"], refs, last)
    base_serial = (base_date - datetime.date(1899, 12, 30)).days

    entries = read_zip(src)
    parts = parts_dict(entries)
    order = [i.filename for i, _ in entries]
    smap = sheet_part_map(parts)
    styles_xml, st = prepare_styles(parts["xl/styles.xml"].decode("utf8"))

    def make(cache):
        newp = dict(parts)
        neworder = list(order)
        newp["xl/styles.xml"] = styles_xml.encode("utf8")
        dh_xml, dh_last = build_sheet("hyg", parts[smap[DH_SHEET][0]].decode("utf8"), cx, st, base_serial, cache)
        dr_xml, dr_last = build_sheet("dr", parts[smap[DR_SHEET][0]].decode("utf8"), cx, st, base_serial, cache)
        newp[smap[DH_SHEET][0]] = dh_xml.encode("utf8")
        newp[smap[DR_SHEET][0]] = dr_xml.encode("utf8")
        newp["xl/workbook.xml"] = patch_workbook(parts["xl/workbook.xml"].decode("utf8"), smap[DR_SHEET][1],
                                                 smap[DH_SHEET][1], dr_last, dh_last, cx.last_slot_row).encode("utf8")
        drop_calc_chain(newp, neworder)
        return newp, neworder

    def write(newp, neworder, path):
        zinfo = {i.filename: i for i, _ in entries}
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            for name in neworder:
                zi = zipfile.ZipInfo(name, date_time=zinfo[name].date_time)
                zi.compress_type = zipfile.ZIP_DEFLATED
                zi.external_attr = zinfo[name].external_attr
                zf.writestr(zi, newp[name])

    tmp = tempfile.mkdtemp(prefix="visit_extract_")
    try:
        p1, o1 = make(None)
        first = os.path.join(tmp, "stage1.xlsx")
        write(p1, o1, first)
        cache = None
        if not no_lo:
            lo_dir = os.path.join(tmp, "lo")
            os.makedirs(lo_dir)
            res = lo_recalc(first, lo_dir)
            if res:
                cache = read_cached_values(res, [DH_SHEET, DR_SHEET])
            else:
                print("[warn] 計算値を埋め込めませんでした(Excelで開くと自動計算されます)。", file=sys.stderr)
        p2, o2 = make(cache)
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        write(p2, o2, out)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return {"out": out, "last_patient_row": last, "slot_rows": [4, cx.last_slot_row],
            "columns": info["columns"], "master": refs, "cached": bool(cache)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("inspect")
    a.add_argument("--src", required=True)
    b = sub.add_parser("build")
    b.add_argument("--src", required=True)
    b.add_argument("--base-date", required=True, help="抽出基準日 YYYY-MM-DD")
    b.add_argument("--out", required=True)
    b.add_argument("--allow-manual", action="store_true", help="手入力セルの扱いを確認済みの場合のみ")
    b.add_argument("--no-lo", action="store_true", help="LibreOfficeでの計算値埋め込みを省略")
    args = ap.parse_args()
    try:
        if args.cmd == "inspect":
            print(json.dumps(inspect_source(args.src), ensure_ascii=False, indent=2))
        else:
            if os.path.abspath(args.src) == os.path.abspath(args.out):
                raise Stop("元ファイルは上書きしません。--out は別のパスにしてください。")
            if os.path.splitext(args.src)[1].lower() != os.path.splitext(args.out)[1].lower():
                raise Stop("--out の拡張子は元ファイルと同じにしてください(.xlsm のマクロを保持するため)。")
            try:
                bd = datetime.date.fromisoformat(args.base_date)
            except ValueError:
                raise Stop("--base-date は YYYY-MM-DD 形式で指定してください。")
            print(json.dumps(run_build(args.src, bd, args.out, args.allow_manual, args.no_lo),
                             ensure_ascii=False, indent=2))
    except Stop as e:
        print(f"[STOP] {e}", file=sys.stderr)
        sys.exit(3)


if __name__ == "__main__":
    main()
