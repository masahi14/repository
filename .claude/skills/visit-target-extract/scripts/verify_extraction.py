#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""抽出結果の検証。数式とは別に書いた「独立モデル」で結果を照合する。

  verify   : 実データの検証(部品の不変確認 / 数式の再計算 / 独立モデルとの照合 / 要確認の一覧)
  selftest : 境界値とランダムデータで、数式がモデルと一致するか確認(複数の基準日)

使い方
  python verify_extraction.py verify   --src 元.xlsx --out 出力.xlsx --base-date 2026-11-01
  python verify_extraction.py selftest --out 出力.xlsx
"""
import argparse
import calendar
import datetime
import json
import os
import random
import re
import shutil
import sys
import tempfile
import warnings
import zipfile

warnings.simplefilter("ignore")
import openpyxl  # noqa: E402
from lxml import etree  # noqa: E402
from openpyxl.utils.datetime import from_excel, to_excel  # noqa: E402

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_extraction as B  # noqa: E402

DOC_NONE = "ドクター診察・治療なし"
REGULAR = "定期検診"
HYG_LIST = ["初診", "定期検診", "SRP", "再評価", "口腔機能", "相談", "それ以外"]
DOC_LIST = ["ドクター診察・治療なし", "初診", "義歯", "充填", "抜歯", "糸抜き", "相談", "その他"]
FREQ_DAYS = {"月1回": 30, "隔週": 14, "週1": 7, "2か月1回": 60, "3か月1回": 90}


# ------------------------------------------------------------ 独立モデル ----
def edate(d, months):
    """Excel の EDATE 相当(暦の月単位。同じ日がなければ月末)。"""
    y, m = divmod(d.year * 12 + d.month - 1 + months, 12)
    m += 1
    return datetime.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def as_date(v):
    """セル値を (kind, date) にする。kind: blank / invalid / date。"""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "blank", None
    if isinstance(v, datetime.datetime):
        return "date", v.date()
    if isinstance(v, datetime.date):
        return "date", v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return "date", from_excel(v).date()
        except Exception:  # noqa: BLE001
            return "invalid", None
    return "invalid", None


def date_state(v, base):
    kind, d = as_date(v)
    if kind != "date":
        return kind
    if d < datetime.date(*B.DATE_MIN) or d > datetime.date(*B.DATE_MAX):
        return "invalid"
    return "future" if d > base else "ok"


def is_future_number(v, base):
    kind, d = as_date(v)
    return kind == "date" and d > base


def tr(v):
    return "" if v is None else str(v).strip()


def model_hyg(p, base):
    """衛生士のみで訪問できる患者の判定。戻り値 (区分 'T'/'C'/None, 理由コード集合)。

    不明な項目は取り得る値をすべて試し、結果が1通りなら確定、割れたら要確認とする。"""
    if p["id"] in (None, "") or tr(p["status"]) == B.ENDED:
        return None, set()
    codes = set()
    st = date_state(p["docLast"], base)
    if st != "ok":
        codes.add("DOC_" + st)
    c = tr(p["ins"])
    ins_poss = {"U"} if c == "利用中" else {"N"} if c == "未利用" else {"U", "N"}
    n = tr(p["hygPlan"])
    plan_poss = ({"T", "O"} if n == "" else {"T"} if n == REGULAR else {"O"} if n in HYG_LIST else {"T", "O"})
    months = set()
    for i in ins_poss:
        if i == "U":
            months.add(3)
        else:
            months |= {2 if q == "T" else 1 for q in plan_poss}
    target = None
    uncertain = False
    if st == "ok":
        d = as_date(p["docLast"])[1]
        outs = {base <= edate(d, m) for m in months}
        uncertain = len(outs) == 2
        target = outs == {True}
    if len(ins_poss) == 2 and (st != "ok" or uncertain):
        codes.add("INS")
    if c != "利用中" and len(plan_poss) == 2 and (st != "ok" or uncertain):
        codes.add("PLAN")
    if is_future_number(p["hygLast"], base):
        codes.add("HYG_FUTURE")
    if is_future_number(p["exam"], base):
        codes.add("EXAM_FUTURE")
    if p["dup"]:
        codes.add("DUP")
    if codes or uncertain:
        return "C", codes
    return ("T" if target else None), codes


def model_dr(p, base):
    if p["id"] in (None, "") or tr(p["status"]) == B.ENDED:
        return None, set()
    codes = set()
    k = tr(p["docPlan"])
    c1 = {"Y"} if k == DOC_NONE else {"N"} if k in DOC_LIST else {"Y", "N"}
    es = date_state(p["exam"], base)
    if es == "ok":
        d = as_date(p["exam"])[1]
        c2 = {"T"} if base < edate(d, B.EXAM_MONTHS) else {"F"}
    else:
        c2 = {"T", "F"}
    outcomes = {(a == "Y" and b == "T") for a in c1 for b in c2}   # True=除外
    uncertain = len(outcomes) == 2
    if uncertain:
        if len(c1) == 2:
            codes.add("K")
        if len(c2) == 2 and es != "future":
            codes.add("EXAM_" + es)
    if es == "future":
        codes.add("EXAM_FUTURE")
    prog = any(p[x] is not None and str(p[x]) != "" for x in ("cr", "ext", "den"))
    if c1 == {"Y"} and prog:
        codes.add("PROGRESS")
    if p["dup"]:
        codes.add("DUP")
    if uncertain or codes:
        return "C", codes
    return (None if outcomes == {True} else "T"), codes


REASON_KEYS = {
    "DOC_blank": "ドクター最終診療日が空欄", "DOC_invalid": "ドクター最終診療日（", "DOC_future": "ドクター最終診療日（",
    "INS": "介護保険の利用状況が空欄または不明", "PLAN": "衛生士次回予定が空欄、または",
    "HYG_FUTURE": "衛生士最終診療日（", "EXAM_FUTURE": "口腔機能低下症の最終検査日（",
    "DUP": "患者IDが重複", "K": "ドクター次回予定", "EXAM_blank": "口腔機能低下症の最終検査日が空欄",
    "EXAM_invalid": "口腔機能低下症の最終検査日（", "PROGRESS": "進捗の記録があります",
}


def reason_ok(codes, text):
    """理由テキストに、期待する理由が含まれ、期待しない理由が混ざっていないか(キーワードで確認)。"""
    t = text or ""
    for c in codes:
        if REASON_KEYS[c] not in t:
            return f"理由「{c}」が表示されていない"
    distinct = {
        "DOC_blank": "ドクター最終診療日が空欄", "INS": "介護保険の利用状況", "PLAN": "衛生士次回予定が空欄",
        "HYG_FUTURE": "衛生士最終診療日（", "DUP": "患者IDが重複", "PROGRESS": "進捗の記録があります",
        "EXAM_blank": "口腔機能低下症の最終検査日が空欄",
    }
    for c, key in distinct.items():
        if c not in codes and key in t:
            return f"想定外の理由「{c}」が表示されている"
    if "DOC_invalid" not in codes and "DOC_future" not in codes and "ドクター最終診療日（" in t:
        return "想定外のドクター最終診療日の理由が表示されている"
    if not (codes & {"EXAM_FUTURE", "EXAM_invalid"}) and "口腔機能低下症の最終検査日（" in t:
        return "想定外の検査日の理由が表示されている"
    return None


# --------------------------------------------------------- 表示値の比較 ----
def nv(x):
    """比較用に値を正規化する。"""
    if x is None or x == "":
        return ""
    if isinstance(x, (datetime.datetime, datetime.date)):
        return float(to_excel(x))
    if isinstance(x, bool):
        return x
    if isinstance(x, (int, float)):
        return float(x)
    return str(x)


def read_patients(ws, cols, last=301):
    people = []
    ids = [ws[f"{cols['id']}{r}"].value for r in range(2, last + 1)]
    for r in range(2, last + 1):
        g = lambda k: ws[f"{cols[k]}{r}"].value  # noqa: E731
        p = {k: g(k) for k in ("id", "name", "ins", "fac", "next", "docLast", "docTx", "docPlan", "hygLast",
                               "hygPlan", "status", "cr", "ext", "den", "srp", "exam", "other")}
        p["row"] = r
        p["freq"] = ws[f"{cols['freq']}{r}"].value if "freq" in cols else None
        p["dup"] = p["id"] not in (None, "") and ids.count(p["id"]) > 1
        people.append(p)
    return people


def expected_lists(people, base):
    out = {"hyg": {"T": [], "C": []}, "dr": {"T": [], "C": []}}
    for p in people:
        for key, fn in (("hyg", model_hyg), ("dr", model_dr)):
            t, codes = fn(p, base)
            if t:
                out[key][t].append((p, codes))
    return out


def dr_next(p):
    f, d = tr(p["freq"]), as_date(p["docLast"])
    if f == "不定期":
        return "未定"
    if p["docLast"] in (None, ""):
        return "記録なし"
    if f not in FREQ_DAYS or d[0] != "date":
        return "要確認"
    return float(to_excel(d[1])) + FREQ_DAYS[f]      # シリアル値で加算(日付演算だと1900年付近でずれる)


def hyg_row_expect(p):
    return [p["id"], p["name"], p["ins"], p["fac"], p["next"], p["docLast"], p["hygLast"], p["hygPlan"],
            p["srp"], p["exam"], p["other"]]


def dr_row_expect(p):
    return [p["id"], p["name"], p["fac"], p["docLast"], dr_next(p), p["docTx"], p["docPlan"], p["cr"],
            p["ext"], p["den"], p["exam"], p["other"]]


def compare_sheet(ws, kind, exp, errors, nlast):
    """抽出先シートの A4〜 を期待値と照合する。"""
    order = [(p, codes, "T") for p, codes in exp["T"]] + [(p, codes, "C") for p, codes in exp["C"]]
    for i, (p, codes, t) in enumerate(order):
        r = 4 + i
        got = [nv(ws.cell(r, c).value) for c in range(1, nlast + 1)]
        want = [nv(x) for x in (hyg_row_expect(p) if kind == "hyg" else dr_row_expect(p))]
        for c, (g, w) in enumerate(zip(got, want), start=1):
            if g != w:
                # 患者一覧の「次回目安日」等が数式で、openpyxl経由の型差(整数/日付)を吸収する
                if isinstance(g, float) and isinstance(w, float) and abs(g - w) < 1e-9:
                    continue
                errors.append(f"{ws.title} 行{r} 列{c} ID{p['id']}: 表示={g!r} 期待={w!r}")
        reason = ws.cell(r, nlast + 1).value or ""
        if t == "C":
            msg = reason_ok(codes, reason)
            if not reason.startswith("要確認："):
                errors.append(f"{ws.title} 行{r} ID{p['id']}: 要確認の理由が空です")
            elif msg:
                errors.append(f"{ws.title} 行{r} ID{p['id']}: {msg} ({reason})")
        elif reason:
            errors.append(f"{ws.title} 行{r} ID{p['id']}: 確定対象なのに理由が表示されています")
    # 余分な行がないこと(重複・前回結果の残存を含む)
    for r in range(4 + len(order), 304):
        vals = [ws.cell(r, c).value for c in range(1, nlast + 2)]
        if any(v not in (None, "") for v in vals):
            errors.append(f"{ws.title} 行{r}: 想定外の行が残っています {vals[:3]}")
    rows_shown = [ws.cell(4 + i, 17).value for i in range(len(order))]
    if len(rows_shown) != len(set(rows_shown)):
        errors.append(f"{ws.title}: 同じ患者(行)が重複して表示されています")


# --------------------------------------------------------- LibreOffice ----
def strip_cached(src, dst, sheets):
    """抽出先シートの保存済み計算値を除いたコピーを作る(LibreOfficeに再計算させるため)。"""
    entries = B.read_zip(src)
    parts = B.parts_dict(entries)
    smap = B.sheet_part_map(parts)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zf:
        for info, data in entries:
            for s in sheets:
                if info.filename == smap[s][0]:
                    x = data.decode("utf8")
                    x = re.sub(r"(</f>)<v>[^<]*</v>", r"\1", x)
                    x = re.sub(r"(</f>)<v/>", r"\1", x)
                    data = x.encode("utf8")
            zf.writestr(info, data)


def lo_values(path, workdir):
    out = B.lo_recalc(path, workdir)
    if not out:
        raise B.Stop("LibreOfficeで再計算できませんでした。")
    return openpyxl.load_workbook(out, data_only=True)


# ------------------------------------------------------------- verify ----
def integrity(src, out):
    a, b = B.parts_dict(B.read_zip(src)), B.parts_dict(B.read_zip(out))
    smap = B.sheet_part_map(a)
    allowed = {smap[B.DH_SHEET][0], smap[B.DR_SHEET][0], "xl/styles.xml", "xl/workbook.xml",
               "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/calcChain.xml"}
    changed = sorted(n for n in set(a) | set(b) if a.get(n) != b.get(n))
    bad = [n for n in changed if n not in allowed]
    msgs = []
    if bad:
        msgs.append("想定外に変更された部品: " + ", ".join(bad))
    for n in a:
        if n.startswith("xl/worksheets/sheet") and n not in allowed:
            x14a = re.findall(r"<x14:dataValidation[ >]", a[n].decode("utf8"))
            x14b = re.findall(r"<x14:dataValidation[ >]", b[n].decode("utf8"))
            if len(x14a) != len(x14b):
                msgs.append(f"{n}: 入力規則の数が変わりました")
    # styles.xml は追記のみであること
    sa, sb = a["xl/styles.xml"].decode("utf8"), b["xl/styles.xml"].decode("utf8")
    for tag in ("numFmts", "cellStyleXfs", "borders", "cellStyles", "dxfs"):
        ma = re.search(rf"<{tag}\b.*?</{tag}>", sa, re.S)
        mb = re.search(rf"<{tag}\b.*?</{tag}>", sb, re.S)
        if (ma and ma.group(0)) != (mb and mb.group(0)):
            msgs.append(f"styles.xml の {tag} が変わっています")
    xa = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", re.search(r"<cellXfs.*?</cellXfs>", sa, re.S).group(0), re.S)
    xb = re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", re.search(r"<cellXfs.*?</cellXfs>", sb, re.S).group(0), re.S)
    if xb[:len(xa)] != xa:
        msgs.append("既存のセル書式(cellXfs)が変更されています")
    for n, data in b.items():
        if n.endswith(".xml") or n.endswith(".rels") or n.endswith(".vml"):
            try:
                etree.fromstring(data)
            except Exception as e:  # noqa: BLE001
                msgs.append(f"{n}: XMLとして不正 ({e})")
    return {"changed_parts": changed, "problems": msgs}


def real_verify(src, out, base):
    res = {"errors": [], "notes": []}
    ig = integrity(src, out)
    res["integrity"] = ig
    res["errors"] += ig["problems"]
    info = B.inspect_source(src)
    cols = info["columns"]
    wsrc = openpyxl.load_workbook(src, data_only=True)[B.SRC_SHEET]
    people = read_patients(wsrc, cols)
    exp = expected_lists(people, base)

    # 1) 保存済みの値 = 独立モデル
    wb_cached = openpyxl.load_workbook(out, data_only=True)
    chk = {}
    for kind, name, nlast in (("hyg", B.DH_SHEET, 11), ("dr", B.DR_SHEET, 12)):
        errs = []
        compare_sheet(wb_cached[name], kind, exp[kind], errs, nlast)
        res["errors"] += [f"[保存値] {e}" for e in errs]
        chk[name] = len(errs)
    # 2) 数式を再計算した値 = 保存済みの値
    tmp = tempfile.mkdtemp(prefix="verify_")
    try:
        stripped = os.path.join(tmp, "stripped.xlsx")
        strip_cached(out, stripped, [B.DH_SHEET, B.DR_SHEET])
        wb_re = lo_values(stripped, os.path.join(tmp, "lo"))
        for kind, name, nlast in (("hyg", B.DH_SHEET, 11), ("dr", B.DR_SHEET, 12)):
            errs = []
            compare_sheet(wb_re[name], kind, exp[kind], errs, nlast)
            res["errors"] += [f"[再計算] {e}" for e in errs]
            for r in range(1, 304):
                for c in range(1, 31):
                    a, b = wb_cached[name].cell(r, c).value, wb_re[name].cell(r, c).value
                    if nv(a) != nv(b):
                        res["errors"].append(f"[再計算差] {name}!{openpyxl.utils.get_column_letter(c)}{r}: {a!r} vs {b!r}")
                        break
        # 数式エラーがないこと
        for name in (B.DH_SHEET, B.DR_SHEET):
            for row in wb_re[name].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("#"):
                        res["errors"].append(f"数式エラー {name}!{cell.coordinate} {cell.value}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # 3) 報告用の集計
    def label(p):
        return f"{tr(p['name'])}(ID {p['id']})"
    res["base_date"] = base.isoformat()
    res["hygienist"] = {"confirmed": len(exp["hyg"]["T"]), "check": len(exp["hyg"]["C"])}
    res["doctor"] = {"confirmed": len(exp["dr"]["T"]), "check": len(exp["dr"]["C"])}
    ws_h, ws_d = wb_cached[B.DH_SHEET], wb_cached[B.DR_SHEET]
    res["check_list"] = {
        "hygienist": [(label(p), ws_h.cell(4 + len(exp["hyg"]["T"]) + i, 12).value)
                      for i, (p, _) in enumerate(exp["hyg"]["C"])],
        "doctor": [(label(p), ws_d.cell(4 + len(exp["dr"]["T"]) + i, 13).value)
                   for i, (p, _) in enumerate(exp["dr"]["C"])],
    }
    # 既存条件(状態=終了)で外れた患者のうち、今回の条件には合致する人
    ended = []
    for p in people:
        if p["id"] not in (None, "") and tr(p["status"]) == B.ENDED:
            q = dict(p, status=None)
            h, _ = model_hyg(q, base)
            d, _ = model_dr(q, base)
            ended.append((label(p), {"衛生士": h, "ドクター": d}))
    res["ended_but_matching"] = ended
    return res


# ----------------------------------------------------------- selftest ----
def rand_people(base, seed, n=250):
    rnd = random.Random(seed)
    people = []
    def near_edate(k):
        d = edate(base, -k) + datetime.timedelta(days=rnd.choice([-1, 0, 0, 1]))
        return datetime.datetime(d.year, d.month, d.day)
    def rand_date(kind):
        x = rnd.random()
        if x < 0.12:
            return None
        if x < 0.17:
            return rnd.choice(["2026/0714", "未記入", "  "])
        if x < 0.20:
            return datetime.datetime(1900, 1, 2)
        if x < 0.27:
            return datetime.datetime.combine(base + datetime.timedelta(days=rnd.randint(1, 20)), datetime.time())
        if x < 0.55:
            return near_edate(rnd.choice([3, B.EXAM_MONTHS] if kind == "exam" else [1, 2, 3]))
        d = base - datetime.timedelta(days=rnd.randint(0, 230 if kind == "exam" else 150))
        return datetime.datetime(d.year, d.month, d.day)
    for i in range(n):
        pid = 1000 + i
        if i in (7, 120):
            pid = 1006                          # 患者IDの重複(i=6の行と合わせて3行が同じID)
        p = {
            "id": pid, "name": f"テスト　患者{i:03d}",
            "ins": rnd.choice(["利用中"] * 4 + ["未利用"] * 3 + [None, "", "利用中 ", "介護あり"]),
            "fac": rnd.choice(["ハートワン潮来", "養護老人ホーム潮来"]),
            "freq": rnd.choice(["月1回", "2か月1回", "3か月1回", "不定期", None]),
            "next": "未定",
            "docLast": rand_date("doc"), "docTx": rnd.choice(["相談", "義歯", None]),
            "docPlan": rnd.choice([DOC_NONE] * 5 + ["相談", "義歯", "その他", None, "", "ドクター診察・治療無し", "抜歯"]),
            "hygLast": rand_date("hyg"),
            "hygPlan": rnd.choice(["定期検診"] * 3 + ["SRP", "再評価", "口腔機能", "相談", "初診", "それ以外", None, "",
                                                       "定期健診", " 定期検診"]),
            "status": rnd.choice([None] * 12 + ["終了"]),
            "cr": rnd.choice([None] * 6 + ["0/2", "1/1"]), "ext": rnd.choice([None] * 12 + ["0/1"]),
            "den": rnd.choice([None] * 12 + ["1/2"]),
            "srp": rnd.choice([None, None, "1/4"]), "exam": rand_date("exam"),
            "other": rnd.choice([None, None, "メモ"]),
        }
        people.append(p)
    return people


NAMED = [
    # 名称, 基準日, 患者値の上書き, 衛生士の期待, ドクターの期待   ("*" はその側を確認しない)
    ("利用中・ドクター最終診療日2026/08/01・3か月以内(11/01は対象)", "2026-11-01",
     dict(ins="利用中", docLast=datetime.datetime(2026, 8, 1)), "T", "*"),
    ("同・11/02は対象外", "2026-11-02",
     dict(ins="利用中", docLast=datetime.datetime(2026, 8, 1)), None, "*"),
    ("未利用・定期検診・2か月(期限当日は対象)", "2026-10-01",
     dict(ins="未利用", hygPlan="定期検診", docLast=datetime.datetime(2026, 8, 1)), "T", "*"),
    ("未利用・定期検診・2か月(翌日は対象外)", "2026-10-02",
     dict(ins="未利用", hygPlan="定期検診", docLast=datetime.datetime(2026, 8, 1)), None, "*"),
    ("未利用・定期検診以外(SRP)・1か月(期限当日は対象)", "2026-09-01",
     dict(ins="未利用", hygPlan="SRP", docLast=datetime.datetime(2026, 8, 1)), "T", "*"),
    ("未利用・定期検診以外(SRP)・1か月(翌日は対象外)", "2026-09-02",
     dict(ins="未利用", hygPlan="SRP", docLast=datetime.datetime(2026, 8, 1)), None, "*"),
    ("月末: 8/31の3か月後は11/30(当日は対象)", "2026-11-30",
     dict(ins="利用中", docLast=datetime.datetime(2026, 8, 31)), "T", "*"),
    ("月末: 8/31の3か月後は11/30(翌日は対象外)", "2026-12-01",
     dict(ins="利用中", docLast=datetime.datetime(2026, 8, 31)), None, "*"),
    ("未利用・衛生士次回予定が空欄・期限が1か月と2か月の間は要確認", "2026-09-15",
     dict(ins="未利用", hygPlan=None, docLast=datetime.datetime(2026, 8, 1)), "C", "*"),
    ("検査2026/05/01・治療なし: 11/01はちょうど6か月経過=除外しない", "2026-11-01",
     dict(docPlan=DOC_NONE, exam=datetime.datetime(2026, 5, 1)), "*", "T"),
    ("検査2026/05/01・治療なし: 10/31は6か月未満=除外", "2026-10-31",
     dict(docPlan=DOC_NONE, exam=datetime.datetime(2026, 5, 1)), "*", None),
    ("治療なし以外・検査から6か月未満: 除外しない(①だけ)", "2026-10-31",
     dict(docPlan="相談", exam=datetime.datetime(2026, 5, 1)), "*", "T"),
    ("治療なし・検査から6か月経過: 除外しない(②だけ違う)", "2026-11-05",
     dict(docPlan=DOC_NONE, exam=datetime.datetime(2026, 5, 1)), "*", "T"),
    ("治療なし・検査日が空欄: 6か月未満と判断せず要確認", "2026-10-31",
     dict(docPlan=DOC_NONE, exam=None), "*", "C"),
    ("次回予定が空欄・検査が6か月未満: 治療なし扱いにせず要確認", "2026-10-31",
     dict(docPlan=None, exam=datetime.datetime(2026, 5, 1)), "*", "C"),
    ("治療なし・検査6か月未満・CRに進捗記録: 要確認", "2026-10-31",
     dict(docPlan=DOC_NONE, exam=datetime.datetime(2026, 5, 1), cr="0/2"), "*", "C"),
    ("検査日が基準日より後: 両シートとも要確認(判定に使わない)", "2026-10-31",
     dict(docPlan="相談", exam=datetime.datetime(2026, 11, 20)), "C", "C"),
    ("状態=終了: 条件に合っても抽出しない(既存条件)", "2026-11-01",
     dict(ins="利用中", docLast=datetime.datetime(2026, 8, 1), status="終了"), None, None),
]


def selftest(out, seeds=(1, 2, 3), bases=("2026-10-31", "2026-11-01", "2026-11-02")):
    errors, ran = [], 0
    info = B.inspect_source(out)
    cols = info["columns"]
    tmp = tempfile.mkdtemp(prefix="selftest_")
    try:
        for bi, bs in enumerate(bases):
            base = datetime.date.fromisoformat(bs)
            for seed in seeds[:1] if bi else seeds[:2]:
                people = rand_people(base, seed * 100 + bi)
                ran += run_case(out, cols, people, base, f"random-{bs}-{seed}", tmp, errors)
        # 名前付きシナリオ(境界値)。基準日ごとにまとめて1回で計算する
        by_base = {}
        for name, bs, ov, eh, ed in NAMED:
            by_base.setdefault(bs, []).append((name, ov, eh, ed))
        for bs, cases in by_base.items():
            base = datetime.date.fromisoformat(bs)
            people = []
            for i, (name, ov, eh, ed) in enumerate(cases):
                p = {"id": 5000 + i, "name": f"境界　{i:02d}", "ins": "利用中", "fac": "ハートワン潮来", "freq": "月1回",
                     "next": "未定", "docLast": datetime.datetime(2026, 1, 1), "docTx": None, "docPlan": "相談",
                     "hygLast": datetime.datetime(2026, 1, 1), "hygPlan": "SRP", "status": None, "cr": None,
                     "ext": None, "den": None, "srp": None, "exam": datetime.datetime(2026, 1, 1), "other": None}
                p.update(ov)
                people.append(p)
            ran += run_case(out, cols, people, base, f"named-{bs}", tmp, errors, named=cases)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return {"cases_run": ran, "errors": errors}


def run_case(out, cols, people, base, tag, tmp, errors, named=None):
    wb = openpyxl.load_workbook(out)          # 数式はそのまま。患者一覧だけ検証用データに差し替える
    ws = wb[B.SRC_SHEET]
    for r in range(2, 302):
        for c in range(1, 23):
            ws.cell(r, c).value = None
    for i, p in enumerate(people):
        r = 2 + i
        for k, v in p.items():
            if k in cols:
                ws[f"{cols[k]}{r}"].value = v
                if isinstance(v, datetime.datetime):
                    ws[f"{cols[k]}{r}"].number_format = "yyyy/mm/dd"
    wb[B.DH_SHEET]["B2"].value = datetime.datetime(base.year, base.month, base.day)
    path = os.path.join(tmp, f"{tag}.xlsx")
    wb.save(path)
    got = lo_values(path, os.path.join(tmp, "lo_" + tag))
    for i, p in enumerate(people):
        p["row"] = 2 + i
        p["dup"] = p["id"] not in (None, "") and sum(1 for q in people if q["id"] == p["id"]) > 1
    exp = expected_lists(people, base)
    errs = []
    for kind, name, nlast in (("hyg", B.DH_SHEET, 11), ("dr", B.DR_SHEET, 12)):
        compare_sheet(got[name], kind, exp[kind], errs, nlast)
    errors += [f"[{tag}] {e}" for e in errs]
    if named:
        for i, (name, ov, eh, ed) in enumerate(named):
            p = people[i]
            for kind, want, fn in (("衛生士", eh, model_hyg), ("ドクター", ed, model_dr)):
                t, _ = fn(p, base)
                if want != "*" and t != want:
                    errors.append(f"[{tag}] シナリオ「{name}」{kind}: モデル={t} 期待={want}")
            # シートの実表示でも確認する
            for sheet, want, nl in ((B.DH_SHEET, eh, 11), (B.DR_SHEET, ed, 12)):
                ids = {got[sheet].cell(r, 1).value: r for r in range(4, 304) if got[sheet].cell(r, 1).value}
                shown = p["id"] in ids
                is_check = shown and bool(got[sheet].cell(ids[p["id"]], nl + 1).value)
                actual = "C" if is_check else ("T" if shown else None)
                if want != "*" and actual != want:
                    errors.append(f"[{tag}] シナリオ「{name}」シート{sheet}: 表示={actual} 期待={want}")
    return 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    v = sub.add_parser("verify")
    v.add_argument("--src", required=True)
    v.add_argument("--out", required=True)
    v.add_argument("--base-date", required=True)
    s = sub.add_parser("selftest")
    s.add_argument("--out", required=True, help="build_extraction.py で作成した抽出済みファイル")
    a = ap.parse_args()
    try:
        if a.cmd == "verify":
            r = real_verify(a.src, a.out, datetime.date.fromisoformat(a.base_date))
        else:
            r = selftest(a.out)
        print(json.dumps(r, ensure_ascii=False, indent=2, default=str))
        sys.exit(1 if r.get("errors") else 0)
    except B.Stop as e:
        print(f"[STOP] {e}", file=sys.stderr)
        sys.exit(3)


if __name__ == "__main__":
    main()
